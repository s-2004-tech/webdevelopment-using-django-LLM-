import fitz
import pandas as pd
from docx import Document
from pptx import Presentation
import pytesseract
from PIL import Image
from transformers import pipeline, BlipProcessor, BlipForQuestionAnswering

# Load models
qa_pipeline = pipeline("question-answering", model="distilbert-base-cased-distilled-squad")
vqa_processor = BlipProcessor.from_pretrained("Salesforce/blip-vqa-base")
vqa_model = BlipForQuestionAnswering.from_pretrained("Salesforce/blip-vqa-base")

def pdf_to_text(path):
    doc = fitz.open(path)
    text = ""
    for page in doc:
        text += page.get_text()
    return text

def docx_to_text(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs])

def pptx_to_text(path):
    prs = Presentation(path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

def txt_to_text(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()

def image_to_text(path):
    image = Image.open(path)
    return pytesseract.image_to_string(image)

def visual_question_answer(path, question):
    image = Image.open(path).convert("RGB")
    inputs = vqa_processor(image, question, return_tensors="pt")
    out = vqa_model.generate(**inputs)
    return vqa_processor.decode(out[0], skip_special_tokens=True)

def answer_text_question(context, question):
    result = qa_pipeline(question=question, context=context)
    return result['answer']

def answer_document_question(path, question):
    path = str(path)

    if path.endswith('.pdf'):
        context = pdf_to_text(path)
        return answer_text_question(context, question)

    elif path.endswith('.docx'):
        context = docx_to_text(path)
        return answer_text_question(context, question)

    elif path.endswith('.pptx'):
        context = pptx_to_text(path)
        return answer_text_question(context, question)

    elif path.endswith('.txt'):
        context = txt_to_text(path)
        return answer_text_question(context, question)

    elif path.endswith('.csv'):
        df = pd.read_csv(path)
        context = df.astype(str).to_string()
        return answer_text_question(context, question)

    elif path.endswith('.xlsx'):
        df = pd.read_excel(path)
        context = df.astype(str).to_string()
        return answer_text_question(context, question)

    elif path.endswith(('.png', '.jpg', '.jpeg', '.PNG')):
        extracted_text = image_to_text(path)
        if len(extracted_text.strip()) > 30:
            return answer_text_question(extracted_text, question)
        else:
            return visual_question_answer(path, question)

    else:
        return "Unsupported file format."
